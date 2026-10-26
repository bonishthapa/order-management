"""CSV / Excel exports. Columns are (key, label, kind) where kind is text|money|qty|date|number."""
import csv
import datetime as dt
import io
from decimal import Decimal

from django.http import HttpResponse

from apps.calendar_utils.fields import bs_or_none


def _value(row, key):
    value = row
    for part in key.split("."):
        value = value.get(part) if isinstance(value, dict) else getattr(value, part, None)
        if callable(value):
            value = value()
    return value


def expand_dates(columns):
    """Every date column is exported twice: AD and the calculated BS value."""
    result = []
    for key, label, kind in columns:
        result.append((key, f"{label} (AD)" if kind == "date" else label, kind))
        if kind == "date":
            result.append((key, f"{label} (BS)", "date_bs"))
    return result


def _cell(row, key, kind):
    value = _value(row, key)
    if value is None or value == "":
        return ""
    if kind == "date_bs":
        try:
            return bs_or_none(value) or ""
        except (ValueError, TypeError):
            return ""
    if kind == "date":
        return value.isoformat() if isinstance(value, (dt.date, dt.datetime)) else str(value)
    if kind in ("money", "qty", "number"):
        try:
            return Decimal(str(value))
        except ArithmeticError:
            return str(value)
    return str(value)


def export_response(rows, columns, filename, fmt="csv", title=None, meta=None):
    columns = expand_dates(columns)
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        ws = wb.active
        ws.title = (title or "Export")[:31]
        start_row = 1
        if title:
            ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=12)
            start_row = 2
            for label, value in (meta or {}).items():
                ws.cell(row=start_row, column=1, value=f"{label}: {value}")
                start_row += 1
            start_row += 1
        for index, (_, label, _) in enumerate(columns, 1):
            ws.cell(row=start_row, column=index, value=label).font = Font(bold=True)
        for r, row in enumerate(rows, start_row + 1):
            for c, (key, _, kind) in enumerate(columns, 1):
                value = _cell(row, key, kind)
                cell = ws.cell(row=r, column=c, value=float(value) if isinstance(value, Decimal) else value)
                if kind == "money":
                    cell.number_format = "#,##0.00"
                elif kind == "qty":
                    cell.number_format = "#,##0.###"
        for index, (_, label, _) in enumerate(columns, 1):
            ws.column_dimensions[ws.cell(row=start_row, column=index).column_letter].width = max(12, min(40, len(label) + 4))
        buffer = io.BytesIO()
        wb.save(buffer)
        response = HttpResponse(buffer.getvalue(),
                                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
        return response
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    if title:
        writer.writerow([title])
        for label, value in (meta or {}).items():
            writer.writerow([f"{label}: {value}"])
        writer.writerow([])
    writer.writerow([label for _, label, _ in columns])
    for row in rows:
        writer.writerow([_cell(row, key, kind) for key, _, kind in columns])
    response = HttpResponse("﻿" + buffer.getvalue(), content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{filename}.csv"'
    return response
