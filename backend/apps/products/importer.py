"""CSV/XLSX product import. All rows are validated first; nothing is saved if any row fails."""
import csv
import io
from decimal import Decimal, InvalidOperation

from django.db import transaction

from apps.core.exceptions import BusinessRuleError
from apps.core.numbering import next_code
from apps.organizations.models import TaxRate, UnitOfMeasure

from .models import Brand, Product, ProductCategory

COLUMNS = ["sku", "name", "barcode", "product_type", "category", "brand", "unit", "purchase_price", "selling_price",
           "tax_rate", "reorder_level", "description"]
REQUIRED = ["name", "unit"]


def read_rows(upload):
    name = upload.name.lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(upload, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(h or "").strip().lower() for h in rows[0]]
        return [{header[i]: ("" if v is None else str(v)).strip() for i, v in enumerate(r) if i < len(header)} for r in rows[1:]
                if any(v not in (None, "") for v in r)]
    if name.endswith(".csv"):
        text = upload.read().decode("utf-8-sig")
        return [{(k or "").strip().lower(): (v or "").strip() for k, v in row.items()} for row in csv.DictReader(io.StringIO(text))]
    raise BusinessRuleError("Upload a .csv or .xlsx file.")


def _decimal(value, field, errors, row_no):
    if value in ("", None):
        return Decimal("0")
    try:
        result = Decimal(str(value).replace(",", ""))
    except InvalidOperation:
        errors.append({"row": row_no, "field": field, "message": f"'{value}' is not a number."})
        return Decimal("0")
    if result < 0:
        errors.append({"row": row_no, "field": field, "message": "Must not be negative."})
    return result


def import_products(organization, user, upload, update_existing=False, dry_run=False):
    rows = read_rows(upload)
    if len(rows) > 5000:
        raise BusinessRuleError("Import at most 5,000 rows at a time.")
    units = {u.code.upper(): u for u in UnitOfMeasure.objects.filter(organization=organization)}
    taxes = {t.name.upper(): t for t in TaxRate.objects.filter(organization=organization)}
    taxes.update({str(t.rate.normalize()): t for t in TaxRate.objects.filter(organization=organization)})
    errors, prepared, seen = [], [], set()
    for row_no, row in enumerate(rows, start=2):
        for field in REQUIRED:
            if not row.get(field):
                errors.append({"row": row_no, "field": field, "message": "Required."})
        sku = row.get("sku", "").upper()
        if sku:
            if sku in seen:
                errors.append({"row": row_no, "field": "sku", "message": "Duplicate SKU in file."})
            seen.add(sku)
        unit = units.get(row.get("unit", "").upper())
        if row.get("unit") and unit is None:
            errors.append({"row": row_no, "field": "unit", "message": f"Unknown unit '{row.get('unit')}'."})
        tax = None
        if row.get("tax_rate"):
            tax = taxes.get(row["tax_rate"].upper().rstrip("%"))
            if tax is None:
                errors.append({"row": row_no, "field": "tax_rate", "message": f"Unknown tax rate '{row['tax_rate']}'."})
        product_type = (row.get("product_type") or "STOCK").upper()
        if product_type not in ("STOCK", "SERVICE"):
            errors.append({"row": row_no, "field": "product_type", "message": "Use STOCK or SERVICE."})
        existing = Product.objects.filter(organization=organization, sku=sku).first() if sku else None
        if existing and not update_existing:
            errors.append({"row": row_no, "field": "sku", "message": f"SKU {sku} already exists."})
        prepared.append({
            "existing": existing, "sku": sku, "name": row.get("name", "")[:200], "barcode": row.get("barcode", "")[:60],
            "product_type": product_type, "category": row.get("category", ""), "brand": row.get("brand", ""),
            "unit": unit, "tax_rate": tax, "description": row.get("description", ""),
            "purchase_price": _decimal(row.get("purchase_price"), "purchase_price", errors, row_no),
            "selling_price": _decimal(row.get("selling_price"), "selling_price", errors, row_no),
            "reorder_level": _decimal(row.get("reorder_level"), "reorder_level", errors, row_no),
        })
    result = {"rows": len(rows), "errors": errors, "created": 0, "updated": 0, "dry_run": dry_run}
    if errors or dry_run:
        return result
    with transaction.atomic():
        for item in prepared:
            category = ProductCategory.objects.get_or_create(organization=organization, name=item["category"])[0] if item["category"] else None
            brand = Brand.objects.get_or_create(organization=organization, name=item["brand"])[0] if item["brand"] else None
            values = {k: item[k] for k in ("name", "barcode", "product_type", "unit", "tax_rate", "description",
                                           "purchase_price", "selling_price", "reorder_level")}
            values.update(category=category, brand=brand)
            if item["existing"]:
                for key, value in values.items():
                    setattr(item["existing"], key, value)
                item["existing"].save()
                result["updated"] += 1
            else:
                Product.objects.create(organization=organization, created_by=user,
                                       sku=item["sku"] or next_code(organization, "product", "ITM"), **values)
                result["created"] += 1
    return result
